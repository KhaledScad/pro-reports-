"""One-off helper: generates sample Account / QES / FISQ Excel files for
smoke-testing the app. Not part of the running application."""
from openpyxl import Workbook

from app.dummy_data import dummy_rows

OUT = "sample_data"
import os
os.makedirs(OUT, exist_ok=True)

# ---- Account.xlsx --------------------------------------------------------
wb = Workbook()
ws = wb.active
ws.title = "Account"
ws.append(["Username", "Password", "Authorization", "Name", "Job", "Location"])
ws.append(["admin", "Admin@123", "ADMIN", "Khaled", "Administrator", "Abu Dhabi"])
ws.append(["auditor1", "Aud@123", "USER", "Fatima Auditor", "Auditor", "Abu Dhabi"])
ws.append(["sup1", "Sup@123", "SUPERVISOR", "Sara Supervisor", "Project Supervisor", "Abu Dhabi"])
wb.save(f"{OUT}/Account.xlsx")

# ---- QES sample: one row per Frame ID, wide quarter columns --------------
wb = Workbook()
ws = wb.active
ws.title = "QES"
headers = ["FRAME_ID", "ESTABLISHMENT_NAME", "ACTIVITY_TYPE_EN", "ISIC4", "ISIC4_DESC_EN",
           "AUDITOR_NAME_EN", "Auditor_User_Name", "Supervisor_USER_Name"]
for q in ["Q1_2026", "Q2_2026"]:
    headers += [f"{q}_Total_Revenue", f"{q}_Total_Expenses", f"{q}_Operations_Cost", f"{q}_RAW_Material",
                f"{q}_Total_Finished_Goods_Begining_Of_period", f"{q}_Total_Finished_Goods_End_of_period",
                f"{q}_TOTAL_EMPLOYEES", f"{q}_Employee_Salary"]
ws.append(headers)
ws.append([
    "F001", "Al Noor Trading LLC", "Trade", "4610", "Wholesale trade",
    "Fatima Auditor", "auditor1", "sup1",
    # Q1_2026
    500000, 120000, 60000, 30000, 20000, 25000, 12, 180000,
    # Q2_2026
    540000, 130000, 62000, 32000, 25000, 28000, 13, 195000,
])
ws.append([
    "F002", "Gulf Manufacturing Co", "Manufacturing", "2599", "Fabricated metal products",
    "Fatima Auditor", "auditor1", "sup1",
    620000, 200000, 90000, 150000, 40000, 45000, 25, 300000,
    650000, 210000, 95000, 155000, 45000, 50000, 26, 312000,
])
wb.save(f"{OUT}/QES_sample.xlsx")

# ---- FISQ sample -----------------------------------------------------------
wb = Workbook()
ws = wb.active
ws.title = "FISQ"
headers = ["FRAME_ID", "ESTABLISHMENT_NAME", "AUDITOR_NAME_EN", "Auditor_User_Name", "Supervisor_USER_Name"]
for q in ["Q1_2026", "Q2_2026"]:
    headers += [f"{q}_Current_Assets", f"{q}_Non_Current_Assets", f"{q}_Current_Liabilities",
                f"{q}_Non_Current_Liabilities", f"{q}_Equity_In_Form",
                f"{q}_Equity_7_1", f"{q}_Equity_7_2", f"{q}_Equity_7_3",
                f"{q}_Equity_7_4", f"{q}_Equity_7_5", f"{q}_Equity_7_6"]
ws.append(headers)
ws.append([
    "F001", "Al Noor Trading LLC", "Fatima Auditor", "auditor1", "sup1",
    300000, 500000, 150000, 200000, 450000, 100000, 200000, 150000, 0, 0, 0,
    320000, 510000, 155000, 205000, 470000, 100000, 220000, 150000, 0, 0, 0,
])
wb.save(f"{OUT}/FISQ_sample.xlsx")

# ===========================================================================
# Dummy: 20 records x 4 quarters in the exact 20-10 upload layouts
# (app/survey_layouts.py), Right_ blocks and all. app/dummy_data.py lists the
# records that are deliberately off so the red checks have something to show.
for survey_type in ("QES", "FISQ"):
    headers, rows = dummy_rows(survey_type, 20)
    wb = Workbook()
    ws = wb.active
    ws.title = survey_type
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(f"{OUT}/{survey_type}_dummy_20x4Q.xlsx")

print("Sample files written to", OUT)
print("Dummy files: QES_dummy_20x4Q.xlsx, FISQ_dummy_20x4Q.xlsx (20 Frame IDs each, in the 20-10 upload layouts)")
