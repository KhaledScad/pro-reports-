"""
Generates a reference XLSX documenting every variable/column used by the
QES and FDI calculation + comparison engines in this app, so an admin
knows exactly what to put in the Excel files they upload. Column names
follow the 20-10 variable list (app/survey_layouts.py); example values are
dummy record F101's Q3 2026 figures.

Produces 6 sheets: QES Calculations, QES Comparisons, FDI Calculations,
FDI Comparisons, and the exact QES / FDI upload layouts with dummy data.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.dummy_data import dummy_rows

RED = "D61F2C"
DARK = "222222"
LIGHT_GRAY = "F7F7F8"
SECTION_FILL = "FDE9E9"

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor=RED)
SECTION_FONT = Font(bold=True, color=DARK, size=11)
SECTION_FILL_STYLE = PatternFill("solid", fgColor=SECTION_FILL)
BORDER = Border(*(Side(style="thin", color="D9D9D9") for _ in range(4)))
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")

COLUMNS = ["Variable / Column name", "Category", "Required?", "Data type",
           "Applies to (business type)", "Description", "Example value"]
WIDTHS = [34, 20, 12, 10, 24, 55, 16]


def write_sheet(ws, title, rows):
    ws.title = title
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    for c, (header, width) in enumerate(zip(COLUMNS, WIDTHS), start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP_CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 28

    r = 2
    for row in rows:
        if row[0] == "__SECTION__":
            cell = ws.cell(row=r, column=1, value=row[1])
            cell.font = SECTION_FONT
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(COLUMNS))
            for c in range(1, len(COLUMNS) + 1):
                ws.cell(row=r, column=c).fill = SECTION_FILL_STYLE
                ws.cell(row=r, column=c).border = BORDER
            ws.row_dimensions[r].height = 20
            r += 1
            continue
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            cell.alignment = WRAP if c in (1, 6) else Alignment(vertical="top", wrap_text=True)
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        r += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{r - 1}"


def write_layout_sheet(ws, title, survey_type):
    """The exact 20-10 upload layout, column for column, with dummy records."""
    ws.title = title
    headers, rows = dummy_rows(survey_type, 5)
    for c, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = 18
    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value).border = BORDER
    ws.freeze_panes = "B2"


def sec(title):
    return ("__SECTION__", title)


MERGED_NOTE = ("Merged files: QES has three quarter blocks labelled Q2_2026 (Q2_2026_ / Right_Q2_2026_ / Right_Right_Q2_2026_) "
               "holding each establishment's latest records in turn; on upload the tool places each block, record by record, in the "
               "quarter its own YEAR / QUARTER values give (mostly Q2_2026, Q1_2026, Q4_2025). "
               "Left_Right_ / Right_ / Right_Right_ copies of ESTABLISHMENT_NAME and Licence_Number go with their block.")

# ============================================================ QES CALC ===
qes_calc_rows = [
    sec("Identification & assignment — FRAME_ID and ESTABLISHMENT_NAME once per row; the rest once per quarter, e.g. Q3_2026_ISIC4"),
    ["FRAME_ID", "Identification", "Required", "Text", "All", "Unique establishment identifier. One row per Frame ID.", "F101"],
    ["ESTABLISHMENT_NAME", "Identification", "Optional", "Text", "All", "Establishment name shown on the record.", "Establishment 1 LLC"],
    ["Licence_Number", "Identification", "Optional", "Text", "All", "Licence number. Searchable from the License Number box; a search combines every Frame ID sharing it, which the Record card lists.", "CN-1000101"],
    ["YEAR / QUARTER", "Identification", "Required for merged files", "Number / Text", "All", "Year and quarter of the block. " + MERGED_NOTE, "2026 / Q3"],
    ["STRATUM", "Identification", "Optional", "Number", "All", "Sample stratum, shown on the Record card.", "1"],
    ["REGION_ID / ADDRESS", "Identification", "Optional", "Text", "All", "Kept with the record; not used by the tool.", "1 / Building 1, Abu Dhabi"],
    ["ACTIVITY_TYPE_EN / ACTIVITY_TYPE_AR", "Identification", "Optional", "Text", "All", "Business activity in English/Arabic. Used with ISIC4 to classify Trade / Industry / Construction / Transport.", "Trade"],
    ["ISIC4", "Identification", "Optional", "Text/Number", "All", "ISIC Rev.4 activity code. First 2 digits are used as a fallback classifier (45-47 Trade, 10-33 Industry, 41-43 Construction, 49-53 Transport).", "4610"],
    ["ISIC4_DESC_EN / ISIC4_DESC_AR", "Identification", "Optional", "Text", "All", "ISIC4 description text, also scanned for classification keywords.", "Wholesale trade"],
    ["AUDITOR_NAME_EN / AUDITOR_NAME_AR", "Assignment", "Optional", "Text", "All", "Auditor's display name.", "Fatima Auditor"],
    ["ENUMERATOR_NAME_EN / ENUMERATOR_NAME_AR", "Assignment", "Optional", "Text", "All", "Enumerator / Key Account Manager display name.", "Ahmed Enumerator"],
    ["Auditor_User_Name", "Assignment", "Required for Auditor role", "Text", "All", "Login username of the Auditor. A user whose Job = 'Auditor' can only open records where this matches their username, in any quarter.", "auditor1"],
    ["ENUMERATOR_User_name", "Assignment", "Required for Enumerator role", "Text", "All", "Login username of the Enumerator / Key Account Manager. Same per-quarter matching as Auditor.", "enum1"],
    ["AUD_Supervisor_USER_Name", "Assignment", "Required for Supervisor role", "Text", "All", "Username of the auditor's supervisor. A Supervisor sees records where this (or ENU_Supervisor_User_Name) matches; the Record card shows the supervisor's name from the account list.", "sup1"],
    ["ENU_Supervisor_User_Name", "Assignment", "Optional", "Text", "All", "Username of the enumerator's supervisor; also gives that Supervisor access.", "sup1"],

    sec("Per-quarter financial inputs — prefix each with the quarter, e.g. Q3_2026_Total_Revenue"),
    ["Total_Revenue", "Financial input", "Required", "Number", "All", "Quarter revenue / sales, excluding VAT.", "448000"],
    ["Total_Expenses", "Financial input", "Required", "Number", "All", "Total quarter expenses. Used for Profit, and as part of Intermediate Consumption for every activity except Industry.", "123200"],
    ["Operations_Cost", "Financial input", "Required", "Number", "All", "Operating / administrative expenses, part of Intermediate Consumption.", "61600"],
    ["RAW_Material", "Financial input", "Required for Trade & Industry", "Number", "Trade, Industry", "Raw materials used in production (Industry: part of Intermediate Consumption) or net purchases of goods for resale (Trade: comes off Total Production).", "31360"],
    ["Total_Finished_Goods_Begining_Of_period", "Financial input", "Required for Trade & Industry", "Number", "Trade, Industry", "Opening inventory value for the quarter.", "19000"],
    ["Total_Finished_Goods_End_of_period", "Financial input", "Required for Trade & Industry", "Number", "Trade, Industry", "Closing inventory value for the quarter. Compared to next quarter's opening value in the comparison page.", "22400"],
    ["IC", "Financial input (override)", "Optional", "Number", "All", "Intermediate Consumption supplied directly in the file. If present and non-zero it REPLACES the calculated formula below.", ""],
    ["TOTAL_EMPLOYEES", "Financial input", "Required for Average Salary", "Number", "All", "Headcount for the quarter.", "11"],
    ["Employee_Salary", "Financial input", "Required for Average Salary", "Number", "All", "Total quarterly employee compensation (used to derive a monthly average). Flagged when it is more than Total_Expenses or more than Total Production.", "168000"],
    ["Banking_Charges", "Financial input", "Required for Industry", "Number", "Industry", "Bank charges. Part of Intermediate Consumption for Industry; also shown in the expense details.", "2800"],
    ["FUND_FOR_RESEARCH", "Other", "Optional", "Number", "All", "Kept with the record; not used by the tool.", "0"],
    ["Comment_Revenue", "Note", "Optional", "Text", "All", "Free-text note on the revenue figure, shown alongside the record.", "Confirmed with client"],
    ["Comment_Expense", "Note", "Optional", "Text", "All", "Free-text note on the expenses figure.", ""],
    ["Comment_Inventory", "Note", "Optional", "Text", "All", "Free-text note on the inventory figures.", ""],

    sec("Calculated outputs (produced by the tool — do not upload these)"),
    ["Business classification", "Calculated", "—", "Text", "All", "Trade / Industry / Construction / Transport / Other, derived from ISIC4 and activity text.", "Trade"],
    ["Total Production", "Calculated", "—", "Number", "All", "Trade: Revenue + End Inv. − Begin Inv. − Raw Material.  Industry: Revenue + End Inv. − Begin Inv.  Construction / Transport / Other: Revenue.", "420,040"],
    ["Intermediate Consumption (IC)", "Calculated", "—", "Number", "All", "The IC column if supplied; otherwise Industry: Operations Cost + Banking Charges + Raw Material, every other activity: Operations Cost + Expenses.", "184,800"],
    ["Added Value", "Calculated", "—", "Number", "All", "Total Production − Intermediate Consumption.", "235,240"],
    ["Profit", "Calculated", "—", "Number", "All", "Trade: Total Production − (Expenses − Raw Material).  Industry: Total Production − Expenses.  Construction / Transport / Other: Revenue − Expenses.", "328,200"],
    ["Average Salary", "Calculated", "—", "Number", "All", "Employee_Salary ÷ TOTAL_EMPLOYEES ÷ 3 (quarterly salary spread over 3 months).", "5,091"],
    ["Profit %", "Calculated", "—", "Percent", "All", "Profit ÷ |Revenue| × 100.", "73.3%"],
    ["Added Value %", "Calculated", "—", "Percent", "All", "Added Value ÷ |Total Production| × 100.", "56.0%"],
    ["Intermediate Consumption %", "Calculated", "—", "Percent", "All", "Intermediate Consumption ÷ |Total Production| × 100.", "44.0%"],
    ["Validation status", "Calculated", "—", "Text", "All", "'Positive added value' / 'Negative added value' / 'Zero added value'.", "Positive added value"],
]

# ============================================================ QES COMP ===
qes_comp_rows = [
    sec("How comparison works"),
    ["Frame ID / License Number", "Selection", "Required", "Text", "All", "Pick a Frame ID or a License Number. Every Frame ID sharing that licence number is combined, each quarter from the Frame ID that reported it, and compared across 2 or more quarters.", "F101"],
    ["Selected periods", "Selection", "Required", "List", "All", "Every quarter present for that Frame ID is selectable (after merged Right_ blocks are renamed on upload); pick 2+ to compare.", "Q2_2026, Q3_2026"],

    sec("Table 1 — Financial comparison (one row per selected quarter)"),
    ["Revenue", "Financial", "Calculated", "Number", "All", "= Total_Revenue for that quarter.", "380,000 / 448,000"],
    ["Expenditure", "Financial", "Calculated", "Number", "All", "= Total_Expenses for that quarter.", "104,500 / 123,200"],
    ["Profit", "Financial", "Calculated", "Number", "All", "Same rule as the single-quarter calculation.", "272,900 / 328,200"],
    ["Added value", "Financial", "Calculated", "Number", "All", "= Total Production − Intermediate Consumption.", "194,050 / 235,240"],
    ["IC", "Financial", "Calculated", "Number", "All", "Intermediate Consumption for that quarter.", "156,750 / 184,800"],
    ["Profit %, AV %, IC %", "Financial", "Calculated", "Percent", "All", "Same ratios as the single-quarter calculation page.", "73.3% / 56.0% / 44.0%"],

    sec("Table 2 — Inventory continuity (one row per selected quarter)"),
    ["Inventory start", "Inventory", "Financial input", "Number", "Trade, Industry", "= Total_Finished_Goods_Begining_Of_period for that quarter.", "19,000"],
    ["Inventory end", "Inventory", "Financial input", "Number", "Trade, Industry", "= Total_Finished_Goods_End_of_period for that quarter.", "22,400"],
    ["Validation", "Inventory", "Calculated", "Text", "Trade, Industry", "Flags whether this quarter's Inventory start equals the PREVIOUS selected quarter's Inventory end.", "Start equals previous end"],

    sec("Table 3 — Employee comparison (one row per selected quarter)"),
    ["Employees", "Employee", "Financial input", "Number", "All", "= TOTAL_EMPLOYEES for that quarter.", "10 / 11"],
    ["Employee compensation", "Employee", "Financial input", "Number", "All", "= Employee_Salary for that quarter.", "142,500 / 168,000"],
    ["Average salary", "Employee", "Calculated", "Number", "All", "= Employee_Salary ÷ Employees ÷ 3.", "5,091"],
    ["Employee change count", "Employee", "Calculated", "Number", "All", "= Employees(this quarter) − Employees(previous quarter). First row shows 'Base period'.", "+1"],
    ["Employee change %", "Employee", "Calculated", "Percent", "All", "= Employee change count ÷ |previous quarter Employees| × 100.", "10.0%"],

    sec("Table 4 — ISIC classification stability"),
    ["ISIC4 per quarter", "Classification", "Financial input", "Text", "All", "Each quarter's ISIC4; flagged red when it changes from the previous quarter. Also cross-checked against the annual AES file's {YYYY}_ISIC4.", "4610"],
]

# ============================================================ FDI CALC ==
fdi_calc_rows = [
    sec("Identification & assignment — FRAMEID, ESTABLISHMENT_NAME and LICENSE_NUMBER once per row; the rest once per quarter, e.g. Q3_2026_STRATUM"),
    ["FRAMEID", "Identification", "Required", "Text", "All", "Unique establishment identifier. One row per Frame ID. Right_FRAMEID (the merged copy) is ignored.", "F101"],
    ["ESTABLISHMENT_NAME", "Identification", "Optional", "Text", "All", "Establishment name. Right_ESTABLISHMENT_NAME is the second quarter's copy.", "Establishment 1 LLC"],
    ["LICENSE_NUMBER", "Identification", "Optional", "Text", "All", "Licence number. Searchable from the License Number box; a search combines every Frame ID sharing it, which the Record card lists. Right_LICENSE_NUMBER is the second quarter's copy.", "CN-1000101"],
    ["STRATUM", "Identification", "Optional", "Number", "All", "Sample stratum, shown on the Record card.", "1"],
    ["VISIT_STATUS_EN / VISIT_STATUS_AR", "Identification", "Optional", "Text", "All", "Outcome of the quarter's visit, shown on the Record card.", "Active"],
    ["ISIC4", "Identification", "Optional", "Text/Number", "All", "ISIC Rev.4 activity code, shown on the Record card; feeds the ISIC stability check.", "4610"],
    ["ISIC4_DESC_EN / ISIC4_DESC_AR", "Identification", "Optional", "Text", "All", "ISIC4 description; kept with the record.", "Wholesale trade"],
    ["AUDITOR_NAME", "Assignment", "Optional", "Text", "All", "Auditor's display name (appears twice in each quarter; either copy works).", "Fatima Auditor"],
    ["ENUMERATOR_NAME", "Assignment", "Optional", "Text", "All", "Enumerator / Key Account Manager display name.", "Ahmed Enumerator"],
    ["Auditor_User_Name", "Assignment", "Required for Auditor role", "Text", "All", "Login username of the Auditor; drives row-level access, checked per quarter.", "auditor1"],
    ["ENUMERATOR_User_name", "Assignment", "Required for Enumerator role", "Text", "All", "Login username of the Enumerator / Key Account Manager.", "enum1"],
    ["AUD_Supervisor_USER_Name", "Assignment", "Required for Supervisor role", "Text", "All", "Username of the auditor's supervisor; the Record card shows the supervisor's name from the account list.", "sup1"],
    ["ENU_Supervisor_User_Name", "Assignment", "Optional", "Text", "All", "Username of the enumerator's supervisor; also gives that Supervisor access.", "sup1"],

    sec("Per-quarter asset equation inputs — prefix each with the quarter, e.g. Q3_2026_Total_Equity"),
    ["Total_of_Current_Assets", "Financial input", "Required", "Number", "All", "Total current assets at the end of the quarter.", "313,600"],
    ["Total_of _Non-Current_Assets", "Financial input", "Required", "Number", "All", "Total non-current assets. (The 20-10 layout's Q1 2026 copy has no quarter prefix; the tool attaches it to Q1 2026 by its position.)", "537,600"],
    ["Total_of_Current_Liabilities", "Financial input", "Required", "Number", "All", "Total current liabilities.", "156,800"],
    ["Total_of_Long-Term_Liabilities", "Financial input", "Required", "Number", "All", "Total long-term (non-current) liabilities.", "212,800"],
    ["Total_Equity", "Financial input", "Required", "Number", "All", "Equity as reported on the survey form; the anchor value the calculation is checked against.", "481,600"],
    ["7.2_Paid-up _Capital_Stock", "Financial input", "Required", "Number", "All", "Paid-up capital (form item 7.2).", "288,960"],
    ["7.3_Other_Equity", "Financial input", "Required", "Number", "All", "Other equity (form item 7.3).", "192,640"],
    ["EQ_Q7_7_1", "Financial input", "Optional", "Number", "All", "Equity item 7.1. Together with 7.2 and 7.3 it adds up to Total_Equity; counted as 0 when the file doesn't have it.", "0"],

    sec("Form items (optional, display only) — the 20-10 layout carries them for the two latest quarters"),
    ["CURR_AS_Q3_D / CURR_AS_Q3_E", "Form item", "Optional", "Number", "All", "Current-asset items D and E (form Q3). Part of Total_of_Current_Assets, never added into it.", "109,760 / 156,800"],
    ["NCURR_AS_Q4_D / NCURR_AS_Q4_E", "Form item", "Optional", "Number", "All", "Non-current-asset items D and E (form Q4).", "322,560 / 161,280"],
    ["CURR_LIAB_Q5C / CURR_LIAB_Q5D", "Form item", "Optional", "Number", "All", "Current-liability items C and D (form Q5).", "78,400 / 47,040"],
    ["LNG_LIAB_Q6C / LNG_LIAB_Q6D", "Form item", "Optional", "Number", "All", "Long-term-liability items C and D (form Q6).", "117,040 / 53,200"],
    ["…_PREV (e.g. CURR_AS_Q3_D_PREV)", "Form item", "Optional", "Number", "All", "The same item's previous-period value, shown under it.", "93,100"],
    ["EQ_Q7_TOTAL_PREV / EQ_Q7_7_3_VAL_PREV", "Form item", "Optional", "Number", "All", "Previous-period total equity (7) and other equity (7.3). Shown under Equity in form and 7.3 on the calculation and comparison pages; not part of any check.", "470,000 / 185,000"],

    sec("Calculated outputs (produced by the tool — do not upload these)"),
    ["Total Assets", "Calculated", "—", "Number", "All", "= Total_of_Current_Assets + Total_of_Non-Current_Assets.", "851,200"],
    ["Total Liabilities", "Calculated", "—", "Number", "All", "= Total_of_Current_Liabilities + Total_of_Long-Term_Liabilities.", "369,600"],
    ["Calculated Equity", "Calculated", "—", "Number", "All", "= Total Assets − Total Liabilities.", "481,600"],
    ["Equation Difference", "Calculated", "—", "Number", "All", "= Calculated Equity − Total_Equity. Should be ~0.", "0"],
    ["7.1 + 7.2 + 7.3 Total", "Calculated", "—", "Number", "All", "= EQ_Q7_7_1 + 7.2_Paid-up_Capital_Stock + 7.3_Other_Equity.", "481,600"],
    ["7.1 + 7.2 + 7.3 Difference / Status", "Calculated", "—", "Number/Text", "All", "= (7.1+7.2+7.3) − Total_Equity; status 'Correct' if within 0.01, else 'Flag / review'.", "0 / Correct"],
    ["Status", "Calculated", "—", "Text", "All", "'Balanced' if both the Equation Difference and the 7.1+7.2+7.3 check are within 0.01, else 'Needs review'.", "Balanced"],
]

# ============================================================ FDI COMP ==
fdi_comp_rows = [
    sec("How comparison works"),
    ["Frame ID / License Number", "Selection", "Required", "Text", "All", "Pick a Frame ID or a License Number. Every Frame ID sharing that licence number is combined, each quarter from the Frame ID that reported it, and compared across 2 or more quarters.", "F101"],
    ["Selected periods", "Selection", "Required", "List", "All", "Every quarter present for that Frame ID is selectable; pick 2+ to compare.", "Q2_2026, Q3_2026"],

    sec("Table 1 — Asset equation comparison (one row per selected quarter)"),
    ["Current assets / Non-current assets / Total assets", "Asset equation", "Calculated", "Number", "All", "Same fields as the single-quarter calculation.", "313,600 / 537,600 / 851,200"],
    ["Current liabilities / Non-current liabilities / Total liabilities", "Asset equation", "Calculated", "Number", "All", "Same fields as the single-quarter calculation.", "156,800 / 212,800 / 369,600"],
    ["Calculated equity / Equity in form / Difference", "Asset equation", "Calculated", "Number", "All", "Same fields as the single-quarter calculation.", "481,600 / 481,600 / 0"],

    sec("Table 2 — Beginning vs end of period (only when the file has opening balances)"),
    ["…_Begining_Of_period / …_End_of_period", "Asset equation", "Optional", "Number", "All", "The 20-10 layout has closing totals only, so this table is hidden for it. Files that do carry opening balances get the asset equation at both points and the movement between them.", "—"],

    sec("Table 3 — ISIC classification stability"),
    ["ISIC4 per quarter", "Classification", "Financial input", "Text", "All", "Each quarter's ISIC4; flagged red when it changes. Also cross-checked against the annual FIS file's {YYYY}_ISIC4.", "4610"],

    sec("Table 4 — Equity details (one row per selected quarter)"),
    ["7.1 / 7.2 Paid-up capital / 7.3 Other equity", "Equity details", "Financial input", "Number", "All", "The three equity items for that quarter; 7.3 also shows its previous-period value (EQ_Q7_7_3_VAL_PREV).", "0 / 288,960 / 192,640"],
    ["7.1+7.2+7.3 / Equity in form / Diff / Status", "Equity details", "Calculated", "Number/Text", "All", "Same check as the single-quarter calculation. Equity in form also shows its previous-period value (EQ_Q7_TOTAL_PREV).", "481,600 / 481,600 / 0 / Correct"],

    sec("Table 5 — Validation (one row per selected quarter)"),
    ["Equation difference / Status", "Validation", "Calculated", "Number/Text", "All", "The equation difference and the overall Balanced / Needs review status.", "0 / Balanced"],

    sec("Table 6 — ±30% movement flag"),
    ["Current / Non-current assets, Current / Non-current liabilities, Equity", "Validation", "Calculated", "Percent", "All", "Flags any of the five totals that moved more than 30% since the previous selected quarter.", "+34% (flagged)"],
]

wb = Workbook()
write_sheet(wb.active, "QES Calculations", qes_calc_rows)
write_sheet(wb.create_sheet(), "QES Comparisons", qes_comp_rows)
write_sheet(wb.create_sheet(), "FDI Calculations", fdi_calc_rows)
write_sheet(wb.create_sheet(), "FDI Comparisons", fdi_comp_rows)
write_layout_sheet(wb.create_sheet(), "QES upload layout (dummy)", "QES")
write_layout_sheet(wb.create_sheet(), "FDI upload layout (dummy)", "FDI")

OUT_PATH = "Proactive_Reporting_QES_FDI_Variables.xlsx"
wb.save(OUT_PATH)
print("Saved", OUT_PATH)
