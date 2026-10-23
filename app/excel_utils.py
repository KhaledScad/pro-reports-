"""
Excel ingestion. Ported from the original tool's `v17ReadWorkbook` /
`v17WorksheetToRows` logic: detect the real header row even if the sheet
has title/notes rows above it, by scanning for a Frame ID-looking cell
(survey files) or a Username/Password pair (Account file).
"""
from __future__ import annotations

import io
import re
from collections import Counter

from openpyxl import load_workbook

from .calc_common import normalize_header

FRAME_ID_TOKENS = {
    "frameid", "frame", "id", "establishmentid", "establishmentnumber",
    "رقمالاطار", "الاطار", "رقمالمنشاة", "رقمالمنشأة",
}


def _is_frame_header_cell(value) -> bool:
    n = normalize_header(value)
    if not n:
        return False
    if n in FRAME_ID_TOKENS:
        return True
    return "frameid" in n or "رقمالاطار" in n


def _unique_headers(headers: list) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, h in enumerate(headers):
        base = str(h).strip() if h not in (None, "") else f"Column_{i + 1}"
        if base not in seen:
            seen[base] = 1
            out.append(base)
        else:
            seen[base] += 1
            out.append(f"{base}_{seen[base]}")
    return out


def _find_header_row(grid: list, is_header_cell) -> int | None:
    for i, row in enumerate(grid):
        if any(is_header_cell(c) for c in row):
            return i
    for i, row in enumerate(grid):
        non_empty = sum(1 for c in row if c is not None and str(c).strip() != "")
        if non_empty >= 2:
            return i
    return None


def _sheet_to_rows(ws, is_header_cell=_is_frame_header_cell) -> list[dict]:
    grid = list(ws.iter_rows(values_only=True))
    if not grid:
        return []
    header_index = _find_header_row(grid, is_header_cell)
    if header_index is None:
        return []
    headers = _unique_headers(list(grid[header_index]))
    rows = []
    for line in grid[header_index + 1:]:
        if not any(c is not None and str(c).strip() != "" for c in line):
            continue
        obj = {headers[c]: ("" if (c >= len(line) or line[c] is None) else line[c]) for c in range(len(headers))}
        rows.append(obj)
    return rows


def sheet_names(file_bytes: bytes) -> list[str]:
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def read_workbook(file_bytes: bytes, only_sheet: str | None = None) -> dict[str, list[dict]]:
    """Return {sheet_name: [row_dict, ...]}. If `only_sheet` is given (and
    not the sentinel "__ALL__"), restrict to that one sheet."""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    names = [only_sheet] if only_sheet and only_sheet != "__ALL__" else wb.sheetnames
    sheets = {name: _sheet_to_rows(wb[name]) for name in names if name in wb.sheetnames}
    wb.close()
    return sheets


def get_frame_id(row: dict) -> str:
    for key, value in row.items():
        if _is_frame_header_cell(key):
            return str(value).strip()
    if row:
        first_key = next(iter(row))
        return str(row[first_key]).strip()
    return ""


_MERGE_PREFIX_RE = re.compile(r"^((?:(?:left|right)_)+)", re.IGNORECASE)
_QUARTER_LABEL_RE = re.compile(r"^(Q[1-4])_(\d{4})_(.+)$", re.IGNORECASE)


def _split_header(header: str, all_headers: set[str]) -> tuple[str, str | None, str]:
    """(merge prefix, quarter label, rest) of one column name. The Left_/Right_
    prefixes a table merge adds to clashing names only count as such when the
    name without them carries a quarter label or also exists in the sheet."""
    m = _MERGE_PREFIX_RE.match(header)
    merge = m.group(1) if m else ""
    base = header[len(merge):]
    q = _QUARTER_LABEL_RE.match(base)
    if merge and not q and base not in all_headers:
        merge, base = "", header  # a real name that merely starts with Left_/Right_
    if q:
        return merge, f"{q.group(1).upper()}_{q.group(2)}", q.group(3)
    return merge, None, base


def _cell_quarter(year_value, quarter_value) -> str | None:
    """The Q#_YYYY that one record's YEAR / QUARTER cells name, if any."""
    y = re.search(r"\d{4}", str(year_value or ""))
    s = str(quarter_value or "")
    q = re.search(r"Q\s*([1-4])", s, re.IGNORECASE) or re.fullmatch(r"\s*([1-4])(?:\.0+)?\s*", s)
    return f"Q{q.group(1)}_{y.group(0)}" if y and q else None


class QuarterLayout:
    """Which quarter each column of one sheet belongs to.

    A sheet merged from several quarterly extracts holds one block of
    Q#_YYYY_-prefixed columns per extract. Extracts that shared a label arrive
    as Q2_2026_ / Right_Q2_2026_ / Right_Right_Q2_2026_ (QES 20-10), and those
    blocks hold each establishment's latest records in turn: mostly Q2 2026,
    Q1 2026 and Q4 2025, but e.g. Q4 2025 in the first block for an
    establishment with nothing newer. So such a block is placed record by
    record from that record's own YEAR / QUARTER cells; where those are empty
    it takes a quarter the record has no figures for. A block whose label no
    other block shares (Q3_2026_) is taken as labelled.

    A column without a quarter label that sits inside a block (Licence_Number,
    FDI's CURR_AS_Q3_D, or a merged copy such as Right_ESTABLISHMENT_NAME,
    which belongs with the block after it) goes with that block. Columns
    before the first block (Frame ID, ESTABLISHMENT_NAME) stay plain and serve
    every quarter."""

    def __init__(self, rows: list[dict]):
        self.headers = list(rows[0].keys()) if rows else []
        all_headers = set(self.headers)
        parts = [_split_header(h, all_headers) for h in self.headers]
        self.rest = [rest for _, _, rest in parts]

        ids: dict[tuple[str, str], int] = {}
        self.labelled = [ids.setdefault((merge, label), len(ids)) if label else None for merge, label, _ in parts]
        shared = Counter(label for _, label in ids)
        self.fixed = {b: label for (merge, label), b in ids.items() if not merge and shared[label] == 1}
        self.moving = [b for b in ids.values() if b not in self.fixed]
        self.block_names = {b: f"{merge}{label}_" for (merge, label), b in ids.items()}
        self.prefixed = {b for (merge, _), b in ids.items() if merge}
        cells: dict[int, dict[str, str]] = {b: {} for b in ids.values()}
        for i, b in enumerate(self.labelled):
            if b is not None:
                cells[b].setdefault(normalize_header(self.rest[i]), self.headers[i])
        self._year_quarter = {b: (cells[b].get("year"), cells[b].get("quarter")) for b in self.moving}

        # A moving block's usual quarter: what most of its records say, else its label.
        taken = set(self.fixed.values())
        self.usual: dict[int, str] = {}
        for (_, label), b in ids.items():
            if b in self.fixed:
                continue
            votes = Counter(q for r in rows if (q := self.own_quarter(r, b)))
            q = votes.most_common(1)[0][0] if votes else label
            if q not in taken:
                taken.add(q)
                self.usual[b] = q

        def nearest(i: int, step: int) -> int | None:
            j = i + step
            while 0 <= j < len(parts):
                if self.labelled[j] is not None:
                    return self.labelled[j]
                j += step
            return None

        self.block = list(self.labelled)
        for i, (merge, label, rest) in enumerate(parts):
            if label or _is_frame_header_cell(rest):
                continue
            before, after = nearest(i, -1), nearest(i, 1)
            if merge:  # a merged table's copy belongs with the block that follows it
                self.block[i] = after if after is not None else before
            elif before == after:  # plain column: only when it sits inside one block
                self.block[i] = before
        self._renames: dict[tuple, dict[str, str | None]] = {}

    def own_quarter(self, row: dict, b: int) -> str | None:
        year, quarter = self._year_quarter.get(b, (None, None))
        return _cell_quarter(row.get(year), row.get(quarter)) if year and quarter else None

    def placement(self, row: dict) -> tuple[str | None, ...]:
        """This record's quarter for each moving block, in block order. None:
        no quarter is free for it, e.g. its own YEAR / QUARTER names one the
        record already has."""
        taken = set(self.fixed.values())
        place: dict[int, str | None] = {}
        for b in self.moving:
            q = self.own_quarter(row, b)
            if q:
                place[b] = None if q in taken else q
                taken.add(q)
        for b in self.moving:
            if b not in place:
                place[b] = next((q for q in [self.usual.get(b), *self.usual.values()] if q and q not in taken), None)
                taken.add(place[b])
        return tuple(place[b] for b in self.moving)

    def apply(self, row: dict) -> dict:
        """`row` with its columns placed (the same dict when nothing moves)."""
        rename = self.renames(row)
        if not rename:
            return row
        return {n: v for k, v in row.items() if (n := rename.get(k, k)) is not None}

    def renames(self, row: dict) -> dict[str, str | None]:
        """New name of each column of `row` that moves; None drops the column."""
        key = self.placement(row) if self.moving else ()
        if key not in self._renames:
            self._renames[key] = self._renames_for(dict(zip(self.moving, key)))
        return self._renames[key]

    def _renames_for(self, place: dict[int, str | None]) -> dict[str, str | None]:
        quarter = {**self.fixed, **place}
        rename: dict[str, str | None] = {}
        used = {h for h, b in zip(self.headers, self.labelled) if b is None or b in self.fixed}
        for i, h in enumerate(self.headers):
            b = self.labelled[i]
            if b is None or b in self.fixed:
                continue
            if quarter[b]:
                target = f"{quarter[b]}_{self.rest[i]}"
                used.add(target)
                if target != h:
                    rename[h] = target
            elif b not in self.prefixed:
                rename[h] = None  # unplaced, and its own name would pass for its label's quarter
        for i, h in enumerate(self.headers):
            b = self.block[i]
            if self.labelled[i] is not None or b is None or not quarter.get(b):
                continue
            target = f"{quarter[b]}_{self.rest[i]}"
            if target not in used:
                rename[h] = target
                used.add(target)
        return rename

    def notes(self, rows: list[dict]) -> list[str]:
        """Short notes on the moving blocks, for the upload message."""
        if not self.moving:
            return []
        differ = left_out = 0
        for row in rows:
            own = {b: self.own_quarter(row, b) for b in self.moving}
            differ += any(q and q != self.usual.get(b) for b, q in own.items())
            left_out += any(own[b] and q is None for b, q in zip(self.moving, self.placement(row)))
        names = " / ".join(self.block_names[b] for b in self.moving)
        usual = " / ".join(self.usual[b] for b in self.moving if b in self.usual)
        notes = [f"{names} columns placed by each record's own YEAR / QUARTER (mostly {usual}"
                 + (f"; {differ} records differ)" if differ else ")")]
        if left_out:
            notes.append(f"{left_out} records: a block whose YEAR / QUARTER repeats a quarter already in the record was left out")
        return notes


def normalize_quarter_columns(rows: list[dict]) -> tuple[list[dict], list[str]]:
    """Give each quarter's columns a plain Q#_YYYY_ prefix, record by record
    (see QuarterLayout). `rows` are one sheet's records, all with the same
    columns. Returns the renamed rows and short notes for the upload message."""
    if not rows:
        return rows, []
    layout = QuarterLayout(rows)
    return [layout.apply(row) for row in rows], layout.notes(rows)


def long_rows_to_wide(rows: list[dict]) -> list[dict] | None:
    """The long layout — one row per Frame ID per quarter, with YEAR and
    QUARTER columns and no Q#_YYYY_ prefixes, as a database view usually
    returns it — in the wide layout: each row's columns get the prefix of its
    own quarter (the Frame ID column stays plain), so one Frame ID's rows add
    up to one record like a wide sheet's. None when `rows` are not long."""
    headers = list(rows[0]) if rows else []
    all_headers = set(headers)
    if any(_split_header(h, all_headers)[1] for h in headers):
        return None
    year = next((h for h in headers if normalize_header(h) == "year"), None)
    quarter = next((h for h in headers if normalize_header(h) == "quarter"), None)
    if not (year and quarter):
        return None
    out = []
    for row in rows:
        q = _cell_quarter(row.get(year), row.get(quarter))
        out.append({k if _is_frame_header_cell(k) else f"{q}_{k}": v for k, v in row.items()} if q else row)
    return out


def records_from_rows(rows: list[dict], source: str) -> tuple[list[dict], list[str]]:
    """One sheet's (or one query's) rows as {frame_id, data} records, with
    each quarter's columns under their Q#_YYYY_ prefix, skipping rows without
    a Frame ID. Also returns notes on how quarter blocks were placed."""
    wide = long_rows_to_wide(rows)
    if wide is not None:
        rows, notes = wide, [f"{len(wide)} rows read as one row per quarter (YEAR / QUARTER columns)"]
    else:
        rows, notes = normalize_quarter_columns(rows)
    out = []
    for row in rows:
        if not row:
            continue
        frame_id = get_frame_id(row)
        if not frame_id:
            continue
        data = dict(row)
        data["__sheet"] = source
        out.append({"frame_id": frame_id, "data": data})
    return out, notes


def rows_for_survey(file_bytes: bytes, sheet_name: str | None = None) -> tuple[list[dict], list[str]]:
    """Flatten sheet(s) into a list of {frame_id, data} dicts, skipping rows
    without a Frame ID. Also returns notes on how merged quarter blocks were
    placed (see QuarterLayout)."""
    out = []
    notes = []
    for sname, rows in read_workbook(file_bytes, sheet_name).items():
        records, sheet_notes = records_from_rows(rows, sname)
        out += records
        notes += sheet_notes
    return out, notes


def _is_account_header_cell(value) -> bool:
    n = normalize_header(value)
    return n in ("username", "password")


def parse_account_sheet(file_bytes: bytes, sheet_name: str | None = None) -> list[dict]:
    """Parse the Account sheet used to seed users. If `sheet_name` is not
    given, uses a sheet literally named 'Account' if present, otherwise the
    first sheet that contains Username/Password columns anywhere in it."""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    try:
        if sheet_name and sheet_name != "__ALL__":
            candidates = [sheet_name] if sheet_name in wb.sheetnames else []
        else:
            named = next((n for n in wb.sheetnames if n.strip().lower() == "account"), None)
            candidates = [named] + list(wb.sheetnames) if named else list(wb.sheetnames)

        rows = []
        for name in candidates:
            if name is None:
                continue
            rows = _sheet_to_rows(wb[name], _is_account_header_cell)
            if rows and any(normalize_header(k) in ("username", "password") for k in rows[0]):
                break
        else:
            rows = []
    finally:
        wb.close()

    if not rows:
        raise ValueError("Could not find Username/Password columns in this file. "
                          "Use the sheet-selection option to pick the correct sheet.")

    users = []
    for row in rows:
        norm = {normalize_header(k): v for k, v in row.items()}
        username = str(norm.get("username", "") or "").strip()
        password = str(norm.get("password", "") or "").strip()
        role_raw = str(norm.get("authorization", "") or "").strip().upper()
        role = role_raw if role_raw in ("ADMIN", "SUPERVISOR", "CONTROLLER") else "USER"
        name = str(norm.get("name", "") or username).strip()
        job = str(norm.get("job", "") or "").strip()
        location = str(norm.get("location", "") or "").strip()
        # Optional Employees column (a CONTROLLER's team, as the users export
        # writes it); None when the sheet has no such column, so teams set on
        # the Admin page are kept.
        employees = ([e for e in re.split(r"[,;\n]+", str(norm.get("employees") or "")) if e.strip()]
                     if "employees" in norm else None)
        # Optional Status column (Active / Suspended, as the users export
        # writes it); None without one, so suspensions set on the Admin page are kept.
        suspended = "suspend" in str(norm.get("status") or "").lower() if "status" in norm else None
        if username and password:
            users.append({
                "username": username, "password": password, "authorization": role,
                "name": name, "job": job, "location": location, "employees": employees,
                "suspended": suspended,
            })
    return users
