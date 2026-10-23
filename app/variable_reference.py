"""
Admin-only downloadable XLSX variable reference for QES and FDI.

Unlike `guide.html` (deliberately worded for end users, with no raw
spreadsheet column names), this file is FOR the people preparing the
upload spreadsheets — it lists every technical column name in the 20-10
variable list, the label auditors see for it in the tool, a description,
and whether it feeds the Calculation page, the Comparison page, or both.
Its last sheet shows the exact upload layout with dummy records under it.

Per-quarter columns carry the quarter in front (`Q3_2026_Total_Revenue`);
that convention, and how merged-file prefixes are handled on upload, is
noted once in the sheet header rather than on every row.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .calculations_qes import DETAIL_GROUPS as QES_DETAIL_GROUPS
from .calculations_fdi import COMPONENT_GROUPS as FDI_COMPONENT_GROUPS
from .dummy_data import dummy_rows

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SECTION_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
SECTION_FONT = Font(bold=True, color="991B1B")
WRAP = Alignment(wrap_text=True, vertical="top")

COLUMNS = [
    "Section", "Technical column name(s)", "Label shown to auditors (EN)",
    "Label shown to auditors (AR)", "Description (EN)", "الوصف (AR)", "Used in",
]
COLUMN_WIDTHS = [22, 46, 26, 26, 58, 58, 16]

_NOT_USED_EN = "Kept with the record; not shown or used by the tool."
_NOT_USED_AR = "تُحفظ مع السجل؛ لا تعرضها الأداة ولا تستخدمها."

# Who's-who columns shared by both layouts (calc_common.record_meta(), auth.py).
ASSIGNMENT_ROWS = [
    ("Assignment", "Auditor_User_Name", "—", "—",
     "Login username of the auditor. A user whose Job is Auditor only sees records with their username here, in any quarter.",
     "اسم دخول المدقق. المستخدم الذي وظيفته مدقق يرى فقط السجلات التي تحمل اسم مستخدمه هنا، في أي ربع.", "Both"),
    ("Assignment", "ENUMERATOR_User_name", "—", "—",
     "Login username of the enumerator / Key Account Manager; the same per-quarter matching for that role.",
     "اسم دخول الباحث / مدير الحساب الرئيسي؛ بنفس المطابقة لكل ربع لهذا الدور.", "Both"),
    ("Assignment", "AUD_Supervisor_USER_Name", "Supervisor name", "اسم المشرف",
     "Username of the auditor's supervisor. A Supervisor sees records with their username here or in ENU_Supervisor_User_Name; the Record card shows this account's name from the account list.",
     "اسم مستخدم مشرف المدقق. يرى المشرف السجلات التي تحمل اسم مستخدمه هنا أو في ENU_Supervisor_User_Name؛ وتعرض بطاقة السجل اسم صاحب هذا الحساب من قائمة الحسابات.", "Both"),
    ("Assignment", "ENU_Supervisor_User_Name", "—", "—",
     "Username of the enumerator's supervisor; also gives that Supervisor access to the record.",
     "اسم مستخدم مشرف الباحث؛ يمنح هذا المشرف أيضاً صلاحية الوصول إلى السجل.", "Both"),
]

_ISIC_ROW = ("Identification", "ISIC4", "ISIC4", "ISIC4",
             "Four-digit ISIC classification code, shown on the Record card. Feeds the ISIC classification-stability check on the Comparison page.",
             "رمز التصنيف ISIC4 المكوّن من أربعة أرقام، يظهر في بطاقة السجل. يغذي فحص استقرار تصنيف ISIC4 في صفحة المقارنة.", "Both")
_STRATUM_ROW = ("Identification", "STRATUM", "Stratum", "الطبقة",
                "Sample stratum, shown on the Record card.", "طبقة العينة، تظهر في بطاقة السجل.", "Both")
_LICENCE_DESC_EN = ("Searchable from the License Number box and shown on the Record card. A search for any Frame ID "
                    "combines every Frame ID with the same number into one record, each quarter taken from the Frame ID "
                    "that reported it; the card lists them. The newest quarter's non-empty value is used.")
_LICENCE_DESC_AR = ("يمكن البحث به من خانة رقم الرخصة ويظهر في بطاقة السجل. البحث بأي رقم إطار يدمج كل أرقام الإطار "
                    "التي تحمل الرقم نفسه في سجل واحد، ويؤخذ كل ربع من رقم الإطار الذي أورده؛ وتعرضها البطاقة. "
                    "يُستخدم أحدث رقم غير فارغ من الأرباع.")

QES_IDENTIFICATION_ROWS = [
    ("Identification", "FRAME_ID", "Frame ID", "رقم الإطار",
     "Uniquely identifies the establishment, one row per Frame ID. Used to find every quarter's figures for it, and to search on the Calculation/Comparison pages.",
     "يحدد المنشأة بشكل فريد، صف واحد لكل رقم إطار. يُستخدم لاسترجاع أرقام كل ربع لها، وللبحث في صفحتي الاحتساب والمقارنة.", "Both"),
    ("Identification", "ESTABLISHMENT_NAME (the other quarter blocks' copies: Left_Right_ / Right_ / Right_Right_ESTABLISHMENT_NAME)",
     "Establishment", "المنشأة", "Establishment's display name.", "الاسم المعروض للمنشأة.", "Both"),
    ("Identification", "Licence_Number (the other quarter blocks' copies: Left_Right_ / Right_ / Right_Right_Licence_Number)",
     "License Number", "رقم الرخصة", _LICENCE_DESC_EN, _LICENCE_DESC_AR, "Both"),
    ("Identification", "YEAR, QUARTER", "—", "—",
     "Year and quarter of the block of columns. The three QES blocks labelled Q2_2026 (Q2_2026_ / Right_Q2_2026_ / Right_Right_Q2_2026_) hold each establishment's latest records in turn, so the tool reads these values for every record to tell which quarter each block holds (mostly Q2 2026, Q1 2026, Q4 2025). A block whose label is unique, such as Q3_2026_, is taken as labelled.",
     "السنة والربع لمجموعة الأعمدة. تحمل مجموعات QES الثلاث المسماة Q2_2026 (Q2_2026_ و Right_Q2_2026_ و Right_Right_Q2_2026_) أحدث سجلات كل منشأة بالتتابع، لذا تقرأ الأداة هذه القيم لكل سجل لتحديد الربع الذي تحمله كل مجموعة (غالباً الربع الثاني 2026 والربع الأول 2026 والربع الرابع 2025). أما المجموعة ذات التسمية الفريدة مثل Q3_2026_ فتؤخذ كما سُمّيت.", "Both"),
    _STRATUM_ROW,
    ("Identification", "REGION_ID, ADDRESS", "—", "—", _NOT_USED_EN, _NOT_USED_AR, "—"),
    ("Identification", "ACTIVITY_TYPE_EN, ACTIVITY_TYPE_AR", "Activity type", "نوع النشاط",
     "Free-text business activity as reported. Scanned for activity keywords (trade/industry/construction/transport, EN+AR) as a fallback classification signal alongside the ISIC4 code range.",
     "النشاط التجاري كنص حر كما ورد. يُفحص بحثاً عن كلمات مفتاحية للنشاط (تجارة/صناعة/إنشاءات/نقل، عربي وإنجليزي) كإشارة تصنيف احتياطية إلى جانب نطاق رمز ISIC4.",
     "Both"),
    ("Identification", "ISIC4", "ISIC4", "ISIC4",
     "Four-digit ISIC classification code. The first two digits drive automatic business-activity classification (Trade 45–47, Industry 10–33, Construction 41–43, Transport 49–53, else Other). Also feeds the ISIC classification-stability check on the Comparison page.",
     "رمز التصنيف ISIC4 المكوّن من أربعة أرقام. أول رقمين يحددان تصنيف النشاط تلقائياً (تجارة 45–47، صناعة 10–33، إنشاءات 41–43، نقل 49–53، وإلا أخرى). يغذي أيضاً فحص استقرار تصنيف ISIC4 في صفحة المقارنة.",
     "Both"),
    ("Identification", "ISIC4_DESC_EN, ISIC4_DESC_AR", "ISIC4 description", "وصف ISIC4",
     "Text description of the ISIC4 code. Also scanned for activity keywords as a fallback classification signal.",
     "وصف نصي لرمز ISIC4. يُفحص أيضاً بحثاً عن كلمات مفتاحية للنشاط كإشارة تصنيف احتياطية.", "Both"),
    ("Assignment", "AUDITOR_NAME_EN, AUDITOR_NAME_AR", "Auditor", "المدقق",
     "Auditor's name, shown on the Record card.", "اسم المدقق، يظهر في بطاقة السجل.", "Both"),
    ("Assignment", "ENUMERATOR_NAME_EN, ENUMERATOR_NAME_AR", "Enumerator", "الباحث",
     "Enumerator's name, shown on the Record card.", "اسم الباحث، يظهر في بطاقة السجل.", "Both"),
    *ASSIGNMENT_ROWS,
]

FDI_IDENTIFICATION_ROWS = [
    ("Identification", "FRAMEID (Right_FRAMEID, the merged copy, is ignored)", "Frame ID", "رقم الإطار",
     "Uniquely identifies the establishment, one row per Frame ID. Used to find every quarter's figures for it, and to search on the Calculation/Comparison pages.",
     "يحدد المنشأة بشكل فريد، صف واحد لكل رقم إطار. يُستخدم لاسترجاع أرقام كل ربع لها، وللبحث في صفحتي الاحتساب والمقارنة.", "Both"),
    ("Identification", "ESTABLISHMENT_NAME (second quarter's copy: Right_ESTABLISHMENT_NAME)",
     "Establishment", "المنشأة", "Establishment's display name.", "الاسم المعروض للمنشأة.", "Both"),
    ("Identification", "LICENSE_NUMBER (second quarter's copy: Right_LICENSE_NUMBER)",
     "License Number", "رقم الرخصة", _LICENCE_DESC_EN, _LICENCE_DESC_AR, "Both"),
    _STRATUM_ROW,
    ("Identification", "VISIT_STATUS_EN, VISIT_STATUS_AR", "Visit status", "حالة الزيارة",
     "Outcome of the quarter's visit, shown on the Record card.", "نتيجة زيارة الربع، تظهر في بطاقة السجل.", "Both"),
    _ISIC_ROW,
    ("Identification", "ISIC4_DESC_EN, ISIC4_DESC_AR", "—", "—", _NOT_USED_EN, _NOT_USED_AR, "—"),
    ("Assignment", "AUDITOR_NAME (appears twice in each quarter; either copy works)", "Auditor", "المدقق",
     "Auditor's name, shown on the Record card.", "اسم المدقق، يظهر في بطاقة السجل.", "Both"),
    ("Assignment", "ENUMERATOR_NAME", "Enumerator", "الباحث",
     "Enumerator's name, shown on the Record card.", "اسم الباحث، يظهر في بطاقة السجل.", "Both"),
    *ASSIGNMENT_ROWS,
]

QES_ROWS = [
    ("Revenue, expenses & cost of goods", "Total_Revenue", "Revenue", "الإيرادات",
     "Total revenue for the quarter. Feeds Total production, Profit, Profit %.",
     "إجمالي الإيرادات خلال الربع. يغذي إجمالي الإنتاج، الربح، ونسبة الربح.", "Both"),
    ("Revenue, expenses & cost of goods", "Total_Expenses", "Expenditure", "المصروفات",
     "Total expenses for the quarter. Feeds Profit, and Intermediate consumption for every activity except Industry. For Trade it already includes the net purchases (Raw material) — the Trade Profit formula backs them out again to avoid double-counting.",
     "إجمالي المصروفات خلال الربع. يغذي الربح، والاستهلاك الوسيط لكل الأنشطة عدا الصناعة. بالنسبة للتجارة، يتضمن أصلاً صافي المشتريات (المواد الخام) — وتطرحها معادلة ربح التجارة مرة أخرى لتفادي احتسابها مرتين.", "Both"),
    ("Revenue, expenses & cost of goods", "Operations_Cost", "Operations cost", "تكلفة التشغيل",
     "Feeds Intermediate consumption: Operations cost + Bank charges + Raw material for Industry, Operations cost + Expenditure for every other activity.",
     "يغذي الاستهلاك الوسيط: تكلفة التشغيل + العمولات البنكية + المواد الخام للصناعة، وتكلفة التشغيل + المصروفات لكل نشاط آخر.", "Both"),
    ("Revenue, expenses & cost of goods", "RAW_Material", "Raw material", "المواد الخام",
     "For Industry: raw material used in production — feeds Intermediate consumption. For Trade: net purchases for resale — feeds Total production and Profit. Not used for Construction/Transport/Other.",
     "للصناعة: المواد الخام المستخدمة في الإنتاج — تغذي الاستهلاك الوسيط. للتجارة: صافي المشتريات لغرض إعادة البيع — تغذي إجمالي الإنتاج والربح. لا تُستخدم للإنشاءات/النقل/أخرى.", "Both"),
    ("Revenue, expenses & cost of goods", "Total_Finished_Goods_Begining_Of_period",
     "Inventory start", "المخزون، أول المدة",
     "Combined finished-goods + work-in-progress inventory value at the start of the quarter (Trade & Industry only). Feeds Total production; checked against the previous quarter's Inventory end in the Inventory continuity check.",
     "قيمة مخزون البضاعة التامة وغير التامة الصنع مجتمعة في بداية الربع (للتجارة والصناعة فقط). يغذي إجمالي الإنتاج؛ ويُقارَن بمخزون آخر المدة للربع السابق في فحص استمرارية المخزون.", "Both"),
    ("Revenue, expenses & cost of goods", "Total_Finished_Goods_End_of_period",
     "Inventory end", "المخزون، آخر المدة",
     "Same as Inventory start, at the close of the quarter.",
     "نفس مخزون أول المدة، في نهاية الربع.", "Both"),
    ("Revenue, expenses & cost of goods", "IC", "Intermediate consumption (as reported)", "الاستهلاك الوسيط (كما ورد)",
     "If present and non-zero, used directly instead of the activity's formula (for Industry it always equals Operations cost + Bank charges + Raw material in the 20-10 data).",
     "إذا وُجدت وكانت غير صفرية، تُستخدم مباشرة بدلاً من معادلة النشاط (للصناعة تساوي دائماً تكلفة التشغيل + العمولات البنكية + المواد الخام في بيانات 20-10).", "Both"),
    ("Revenue, expenses & cost of goods", "Banking_Charges", "Banking Charges", "العمولات البنكية",
     "Feeds Intermediate consumption for Industry (Operations cost + Bank charges + Raw material); also shown in the Expense details.",
     "يغذي الاستهلاك الوسيط للصناعة (تكلفة التشغيل + العمولات البنكية + المواد الخام)؛ ويظهر أيضاً في تفاصيل المصروفات.", "Both"),
    ("Revenue, expenses & cost of goods", "FUND_FOR_RESEARCH", "—", "—", _NOT_USED_EN, _NOT_USED_AR, "—"),
    ("Employment", "Employee_Salary", "Total compensation", "إجمالي التعويضات",
     "Total employee compensation for the quarter. Feeds Average compensation (÷ Employees ÷ 3). Flagged when it is more than Expenditure (Total_Expenses), since compensation is part of expenditure, and when it is more than Total production.",
     "إجمالي تعويضات الموظفين خلال الربع. يغذي متوسط التعويضات (÷ عدد الموظفين ÷ 3). يُشار إليه بتنبيه عندما يكون أكبر من المصروفات (Total_Expenses)، لأن التعويضات جزء من المصروفات، وعندما يكون أكبر من إجمالي الإنتاج.", "Both"),
    ("Employment", "TOTAL_EMPLOYEES", "Employees", "عدد الموظفين",
     "Headcount for the quarter. Feeds Average compensation; also shown in the Employee comparison table with the period-over-period change (count and %).",
     "عدد الموظفين خلال الربع. يغذي متوسط التعويضات؛ ويظهر أيضاً في جدول مقارنة الموظفين مع التغيّر بين الفترات (بالعدد والنسبة).", "Both"),
    ("Comments", "Comment_Revenue", "Revenue comment", "ملاحظة الإيرادات",
     "Free-text auditor note shown alongside the revenue figure.", "ملاحظة نصية من المدقق تظهر بجانب رقم الإيرادات.", "Calculation"),
    ("Comments", "Comment_Expense", "Expense comment", "ملاحظة المصروفات",
     "Free-text auditor note shown alongside the expenditure figure.", "ملاحظة نصية من المدقق تظهر بجانب رقم المصروفات.", "Calculation"),
    ("Comments", "Comment_Inventory", "Inventory comment", "ملاحظة المخزون",
     "Free-text auditor note shown alongside the inventory figures.", "ملاحظة نصية من المدقق تظهر بجانب أرقام المخزون.", "Calculation"),
]

FDI_ROWS = [
    ("Assets", "Total_of_Current_Assets", "Current assets", "الأصول المتداولة",
     "Total current assets at the end of the quarter. Feeds Total assets.",
     "إجمالي الأصول المتداولة في نهاية الربع. يغذي إجمالي الأصول.", "Both"),
    ("Assets", "Total_of _Non-Current_Assets", "Non-current assets", "الأصول غير المتداولة",
     "Total non-current assets (property, equipment, ...). Feeds Total assets. In the 20-10 layout the Q1 2026 copy has no quarter prefix; the tool attaches it to Q1 2026 by its position.",
     "إجمالي الأصول غير المتداولة (ممتلكات، معدات، ...). يغذي إجمالي الأصول. في تخطيط 20-10 نسخة الربع الأول 2026 بلا بادئة الربع؛ تنسبها الأداة إلى الربع الأول 2026 حسب موقعها.", "Both"),
    ("Liabilities", "Total_of_Current_Liabilities", "Current liabilities", "المطلوبات المتداولة",
     "Obligations due within the year. Feeds Total liabilities.",
     "الالتزامات المستحقة خلال السنة. يغذي إجمالي المطلوبات.", "Both"),
    ("Liabilities", "Total_of_Long-Term_Liabilities", "Non-current liabilities", "المطلوبات غير المتداولة",
     "Long-term liabilities. Feeds Total liabilities.", "المطلوبات طويلة الأجل. يغذي إجمالي المطلوبات.", "Both"),
    ("Equity", "Total_Equity", "Equity in form", "حقوق الملكية في الاستمارة",
     "Total equity as reported in the form: the figure Calculated equity (Total assets − Total liabilities) is checked against.",
     "إجمالي حقوق الملكية كما ورد في الاستمارة: الرقم الذي تُقارن به حقوق الملكية المحتسبة (إجمالي الأصول − إجمالي المطلوبات).", "Both"),
    ("Equity", "EQ_Q7_TOTAL_PREV", "7 Total equity — previous period", "7 إجمالي حقوق الملكية — الفترة السابقة",
     "Total equity reported for the previous period. Shown under Equity in form on the calculation and comparison pages; not part of any check.",
     "إجمالي حقوق الملكية المُبلّغ عنه للفترة السابقة. يظهر تحت حقوق الملكية في الاستمارة في صفحتي الاحتساب والمقارنة؛ لا يدخل في أي فحص.", "Both"),
    ("Equity", "EQ_Q7_7_1", "7.1 Equity item 7.1", "7.1 بند حقوق الملكية 7.1",
     "Equity item 7.1. Together with 7.2 and 7.3 it should add up to Total_Equity (the 7.1 + 7.2 + 7.3 check).",
     "بند حقوق الملكية 7.1. يجب أن يساوي مع البندين 7.2 و7.3 إجمالي حقوق الملكية (فحص 7.1 + 7.2 + 7.3).", "Both"),
    ("Equity", "7.2_Paid-up _Capital_Stock", "7.2 Paid-up capital", "7.2 رأس المال المدفوع",
     "Paid-up capital. See 7.1.", "رأس المال المدفوع. انظر 7.1.", "Both"),
    ("Equity", "7.3_Other_Equity (or EQ_Q7_7_3_VAL)", "7.3 Other equity", "7.3 حقوق ملكية أخرى",
     "Other equity. See 7.1.", "حقوق ملكية أخرى. انظر 7.1.", "Both"),
    ("Equity", "EQ_Q7_7_3_VAL_PREV", "7.3 Other equity — previous period", "7.3 حقوق ملكية أخرى — الفترة السابقة",
     "Other equity reported for the previous period. Shown under 7.3 on the calculation and comparison pages; not part of any check.",
     "حقوق الملكية الأخرى المُبلّغ عنها للفترة السابقة. تظهر تحت البند 7.3 في صفحتي الاحتساب والمقارنة؛ لا تدخل في أي فحص.", "Both"),
    ("Older column names (still accepted)",
     "Current_Assets, Non_Current_Assets, Current_Liabilities, Non_Current_Liabilities, Equity_In_Form, Equity_7_1, Equity_7_2, Equity_7_3 — and their _End_of_period / _Begining_Of_period forms",
     "—", "—",
     "Read only when the 20-10 name is missing, so files in the earlier layout still work. The _Begining_Of_period forms also switch on the Beginning vs end of period table.",
     "تُقرأ فقط عند غياب اسم 20-10، لتبقى الملفات بالتخطيط السابق صالحة. كما تُفعّل صيغ بداية المدة جدول بداية المدة مقابل نهاية المدة.", "Both"),
]

ANNUAL_ISIC_ROWS = [
    ("Frame_ID (or FRAME, ID, ESTABLISHMENTID, ESTABLISHMENTNUMBER)", "Frame ID", "رقم الإطار",
     "Same Frame ID as the QES/FDI quarterly files — used to match this annual record to the establishment.",
     "نفس رقم الإطار المستخدم في ملفات QES/FDI الربعية — يُستخدم لمطابقة هذا السجل السنوي بالمنشأة.", "Comparison"),
    ("{YYYY}_ISIC4 — one column per year, e.g. 2023_ISIC4, 2024_ISIC4, 2025_ISIC4",
     "(year column, e.g. \"2025\")", "(عمود السنة، مثل \"2025\")",
     "The ISIC4 code reported for that year. The tool reads every {YYYY}_ISIC4 column present, sorts them by year, and shows the last 3 on the Comparison page — flagging red the moment the value changes from one year to the next.",
     "رمز ISIC4 المُبلغ عنه لتلك السنة. تقرأ الأداة كل عمود {YYYY}_ISIC4 موجود، وترتبها حسب السنة، وتعرض آخر 3 سنوات في صفحة المقارنة — وتُشير بالأحمر لحظة تغيّر القيمة من سنة لأخرى.", "Comparison"),
]


def _write_note(ws: Worksheet, text_en: str, text_ar: str, row: int) -> int:
    ws.cell(row=row, column=1, value=text_en).font = Font(italic=True, color="6B7280")
    ws.cell(row=row, column=1).alignment = WRAP
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLUMNS))
    row += 1
    ws.cell(row=row, column=1, value=text_ar).font = Font(italic=True, color="6B7280")
    ws.cell(row=row, column=1).alignment = WRAP
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLUMNS))
    return row + 2


def _write_sheet(ws: Worksheet, rows: list[tuple], note_en: str, note_ar: str) -> None:
    ws.sheet_view.rightToLeft = False
    for i, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    r = _write_note(ws, note_en, note_ar, 1)

    header_row = r
    for c, title in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=c, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
    ws.freeze_panes = f"A{header_row + 1}"
    r += 1

    current_section = None
    for section, tech, label_en, label_ar, desc_en, desc_ar, used_in in rows:
        if section != current_section:
            cell = ws.cell(row=r, column=1, value=section)
            for c in range(1, len(COLUMNS) + 1):
                ws.cell(row=r, column=c).fill = SECTION_FILL
            cell.font = SECTION_FONT
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(COLUMNS))
            current_section = section
            r += 1
        values = ["", tech, label_en, label_ar, desc_en, desc_ar, used_in]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = WRAP
        r += 1


def _write_layout_sheet(ws: Worksheet, survey_type: str) -> None:
    """The exact 20-10 upload layout, column for column, with dummy records."""
    headers, rows = dummy_rows(survey_type, 5)
    ws.cell(row=1, column=1, value=_LAYOUT_NOTE_EN).font = Font(italic=True, color="6B7280")
    ws.cell(row=2, column=1, value=_LAYOUT_NOTE_AR).font = Font(italic=True, color="6B7280")
    for c, title in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(c)].width = 18
    for r, values in enumerate(rows, start=5):
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)
    ws.freeze_panes = "B5"


_PERIOD_NOTE_EN = ("Per-quarter columns carry the quarter in front, e.g. Q3_2026_Total_Revenue — one copy per quarter. "
                    "Files merged from several quarterly extracts may also carry Left_Right_ / Right_ / Right_Right_ prefixes: "
                    "on upload the tool attaches each such column, and any column without a quarter label that sits inside a "
                    "quarter's block, to that quarter. Blocks that share one label (QES: Q2_2026_ / Right_Q2_2026_ / "
                    "Right_Right_Q2_2026_) are placed record by record from their own YEAR / QUARTER values. "
                    "The last sheet shows the exact upload layout with dummy records.")
_PERIOD_NOTE_AR = ("تحمل الأعمدة الربعية الربع في بدايتها، مثل Q3_2026_Total_Revenue — نسخة لكل ربع. "
                    "قد تحمل الملفات المدمجة من عدة مستخرجات ربعية أيضاً البادئات Left_Right_ / Right_ / Right_Right_: "
                    "عند الرفع تنسب الأداة كل عمود منها، وكل عمود بلا تسمية ربع يقع ضمن مجموعة ربع، إلى ذلك الربع. "
                    "أما المجموعات التي تشترك في تسمية واحدة (في QES: Q2_2026_ و Right_Q2_2026_ و Right_Right_Q2_2026_) "
                    "فتُنسب سجلاً سجلاً حسب قيم YEAR / QUARTER الخاصة بها. "
                    "تعرض الورقة الأخيرة تخطيط الرفع بالضبط مع سجلات تجريبية.")

_LAYOUT_NOTE_EN = ("Exact upload layout from \"QES and FDI varibles 20-10.xlsx\", column for column, with 5 dummy records. "
                    "Upload this sheet as it is to try the tool; every figure is made up.")
_LAYOUT_NOTE_AR = ("تخطيط الرفع كما في ملف \"QES and FDI varibles 20-10.xlsx\" عموداً بعمود، مع 5 سجلات تجريبية. "
                    "يمكن رفع هذه الورقة كما هي لتجربة الأداة؛ كل الأرقام وهمية.")

_ANNUAL_NOTE_EN = ("This sheet is not part of the QES/FDI upload file — it's the separate annual AES / FIS "
                    "reference file (Admin panel → AES / FIS upload cards) used only for the ISIC "
                    "classification-stability cross-check on the Comparison page.")
_ANNUAL_NOTE_AR = ("هذه الورقة ليست جزءاً من ملف رفع QES/FDI — بل هي ملف AES / FIS السنوي المرجعي المنفصل "
                    "(لوحة الإدارة ← بطاقتا رفع AES / FIS) المستخدم فقط لمطابقة استقرار تصنيف ISIC4 "
                    "في صفحة المقارنة.")


_DETAIL_DESC_EN = "Optional survey-form line item, shown for reference on the Calculation/Comparison pages — read-only, does not feed any formula above."
_DETAIL_DESC_AR = "بند اختياري من استمارة المسح، يُعرض للمرجعية في صفحتي الاحتساب والمقارنة — للعرض فقط، ولا يدخل في أي معادلة أعلاه."
_FDI_ITEM_DESC_EN = ("Survey-form item, shown with its previous-period value (the _PREV column) for reference. It is part of "
                       "the total above, never added into it.")
_FDI_ITEM_DESC_AR = "بند من استمارة المسح، يُعرض مع قيمته للفترة السابقة (عمود _PREV) للمرجعية. هو جزء من الإجمالي أعلاه، ولا يُضاف إليه."

_DETAIL_NOTE_EN = ("Every column below is OPTIONAL and read-only: if present, it appears in a collapsible "
                    "\"Survey form details\" section on the Calculation/Comparison pages, purely for reference. "
                    "None of these feed the totals on the first sheet — those keep being read from their own "
                    "columns, regardless of whether these columns are supplied.")
_DETAIL_NOTE_AR = ("كل عمود أدناه اختياري وللعرض فقط: إذا وُجد، يظهر في قسم قابل للطي بعنوان \"تفاصيل استمارة المسح\" "
                    "في صفحتي الاحتساب والمقارنة، للمرجعية فقط. لا يدخل أي منها في الإجماليات في الورقة الأولى — "
                    "فتلك تُقرأ من أعمدتها بغض النظر عن توفر هذه الأعمدة.")


def _qes_detail_rows() -> list[tuple]:
    rows = []
    for title_en, title_ar, items in QES_DETAIL_GROUPS:
        for label_en, label_ar, codes in items:
            tech = " + ".join(codes) if len(codes) > 1 else codes[0]
            rows.append((title_en, tech, label_en, label_ar, _DETAIL_DESC_EN, _DETAIL_DESC_AR, "Both"))
    return rows


def _fdi_detail_rows() -> list[tuple]:
    rows = []
    for g in FDI_COMPONENT_GROUPS:
        for label_en, label_ar, codes in g["components"]:
            tech = f"{codes[0]} (previous period: {codes[0]}_PREV)"
            rows.append((g["title_en"], tech, label_en, label_ar, _FDI_ITEM_DESC_EN, _FDI_ITEM_DESC_AR, "Both"))
    return rows


def build_workbook(survey_type: str) -> bytes:
    """survey_type: 'QES' or 'FDI'. Returns the .xlsx file bytes."""
    survey_type = survey_type.upper()
    wb = Workbook()
    ws = wb.active
    ws.title = survey_type

    rows = (QES_IDENTIFICATION_ROWS + QES_ROWS) if survey_type == "QES" else (FDI_IDENTIFICATION_ROWS + FDI_ROWS)
    _write_sheet(ws, rows, _PERIOD_NOTE_EN, _PERIOD_NOTE_AR)

    ws2 = wb.create_sheet("AES-FIS annual ISIC")
    annual_rows = [("Frame ID & annual ISIC4", *row) for row in ANNUAL_ISIC_ROWS]
    _write_sheet(ws2, annual_rows, _ANNUAL_NOTE_EN, _ANNUAL_NOTE_AR)

    ws3 = wb.create_sheet("Survey form details (optional)")
    detail_rows = _qes_detail_rows() if survey_type == "QES" else _fdi_detail_rows()
    _write_sheet(ws3, detail_rows, _DETAIL_NOTE_EN, _DETAIL_NOTE_AR)

    _write_layout_sheet(wb.create_sheet("Upload layout (dummy data)"), survey_type)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
